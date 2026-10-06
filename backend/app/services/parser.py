"""Document parsing service — extracts text from PDF, DOCX, and image files."""

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from PIL import Image

logger = logging.getLogger(__name__)


@dataclass
class PageContent:
    """Text content from a single page or section."""
    page_number: int
    text: str
    char_count: int = 0

    def __post_init__(self) -> None:
        self.char_count = len(self.text)


@dataclass
class ParsedDocument:
    """Result of parsing a document."""
    pages: list[PageContent] = field(default_factory=list)
    total_chars: int = 0
    page_count: int = 0
    metadata: dict = field(default_factory=dict)

    def __post_init__(self) -> None:
        self.total_chars = sum(p.char_count for p in self.pages)
        self.page_count = len(self.pages)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


class DocumentParser:
    """Extracts text from supported document formats."""

    SUPPORTED_TYPES: dict[str, str] = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "image/png": "image",
        "image/jpeg": "image",
        "image/webp": "image",
        "image/bmp": "image",
        "image/tiff": "image",
    }

    def parse(self, file_path: Path, media_type: str) -> ParsedDocument:
        """Parse a document file and return extracted text."""
        fmt = self.SUPPORTED_TYPES.get(media_type)
        if fmt is None:
            raise ValueError(f"Unsupported media type: {media_type}")

        if fmt == "pdf":
            return self._parse_pdf(file_path)
        elif fmt == "docx":
            return self._parse_docx(file_path)
        elif fmt == "image":
            return self._parse_image(file_path)
        else:
            raise ValueError(f"Unknown format: {fmt}")

    def _parse_pdf(self, file_path: Path) -> ParsedDocument:
        """Extract text from PDF using PyMuPDF."""
        pages: list[PageContent] = []
        metadata: dict = {}

        try:
            doc = fitz.open(str(file_path))
            metadata = {
                "title": doc.metadata.get("title", ""),
                "author": doc.metadata.get("author", ""),
                "page_count": len(doc),
            }

            for page_num, page in enumerate(doc, start=1):
                text = page.get_text("text").strip()

                # If page has very little text, try extracting from blocks
                if len(text) < 50:
                    blocks = page.get_text("blocks")
                    block_texts = [
                        b[4].strip()
                        for b in blocks
                        if b[6] == 0 and isinstance(b[4], str)
                    ]
                    alt_text = "\n".join(block_texts).strip()
                    if len(alt_text) > len(text):
                        text = alt_text

                if text:
                    pages.append(PageContent(page_number=page_num, text=text))

            doc.close()
        except Exception:
            logger.exception("Failed to parse PDF: %s", file_path)
            raise

        result = ParsedDocument(pages=pages, metadata=metadata)
        logger.info(
            "Parsed PDF: %d pages, %d chars", result.page_count, result.total_chars
        )
        return result

    def _parse_docx(self, file_path: Path) -> ParsedDocument:
        """Extract text from DOCX using python-docx."""
        try:
            doc = DocxDocument(str(file_path))
        except Exception:
            logger.exception("Failed to parse DOCX: %s", file_path)
            raise

        paragraphs: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Extract table contents
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        # Group paragraphs into logical pages (~3000 chars each)
        pages: list[PageContent] = []
        current_text = ""
        page_num = 1

        for para in paragraphs:
            if len(current_text) + len(para) > 3000 and current_text:
                pages.append(PageContent(page_number=page_num, text=current_text))
                page_num += 1
                current_text = para
            else:
                current_text = f"{current_text}\n{para}" if current_text else para

        if current_text:
            pages.append(PageContent(page_number=page_num, text=current_text))

        metadata = {
            "title": doc.core_properties.title or "",
            "author": doc.core_properties.author or "",
            "page_count": len(pages),
        }

        result = ParsedDocument(pages=pages, metadata=metadata)
        logger.info(
            "Parsed DOCX: %d sections, %d chars",
            result.page_count,
            result.total_chars,
        )
        return result

    def _parse_image(self, file_path: Path) -> ParsedDocument:
        """Extract text from image using OCR (Pytesseract) if available, else return placeholder."""
        text = ""

        try:
            import pytesseract

            image = Image.open(file_path)
            text = pytesseract.image_to_string(image).strip()
            logger.info("OCR extracted %d chars from image", len(text))
        except ImportError:
            logger.warning(
                "pytesseract not installed — image will be processed by LLM vision"
            )
            # Return empty text; the summarizer will use Gemini Vision
            text = ""
        except Exception:
            logger.exception("OCR failed for image: %s", file_path)
            text = ""

        pages = []
        if text:
            pages.append(PageContent(page_number=1, text=text))

        return ParsedDocument(
            pages=pages,
            metadata={"type": "image", "page_count": 1},
        )


# Module-level singleton
document_parser = DocumentParser()
